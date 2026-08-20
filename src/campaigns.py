"""Campaign persistence — one folder per campaign on disk.

Layout:
  data/campaigns/<campaign_id>/
    meta.json
    rules_summary.json     # structured brief: criteria / safety / style / submission
    rules_full.<ext>
    template.json          # optional Style Lab draft
    clips.json
    input/  output/  output/raw/
    transcripts/  context/  clip_candidates/  frames/
"""
from __future__ import annotations

import json
import re
import uuid
from datetime import datetime, timezone
from pathlib import Path

from src.config import config
from src.llm_client import call_ollama

CAMPAIGN_STATUSES = ("active", "submitted", "paid", "expired")
CLIP_STATUSES = ("analyzing", "reviewing", "exported", "posted")
RULES_EXTS = {".pdf", ".docx", ".txt", ".md"}
SAFE_ID = re.compile(r"^[a-zA-Z0-9][a-zA-Z0-9._-]{0,80}$")
RULE_LIST_SECTIONS = ("content_criteria", "brand_safety", "editing_style")
RULE_SECTIONS = RULE_LIST_SECTIONS + ("submission_requirements",)

EMPTY_RULES = {
    "content_criteria": [],
    "brand_safety": [],
    "editing_style": [],
    "submission_requirements": "",
    "submission_done": False,
}

SUMMARIZE_SYSTEM = (
    "You condense a creator / clipping brief into structured JSON with "
    "exactly these keys:\n"
    "- content_criteria: string array. What to look for or avoid when "
    "picking clips (hooks, framing, topics). Short bullets.\n"
    "- brand_safety: string array. HARD non-negotiable constraints only "
    "(banned content, required exclusions, 'THIS IS VERY IMPORTANT' "
    "items). Filters, not preferences.\n"
    "- editing_style: string array. Tone, pacing, caption, music, visual "
    "preferences.\n"
    "- submission_requirements: one string. Posting/reporting obligations "
    "as written (analytics forms, view thresholds, audience geo, "
    "engagement minimums, deadlines, links). Empty string if none.\n"
    "Keep each list to the concrete items in the brief. No preamble."
)


def _now():
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def _read_json(path: Path, default=None):
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return default


def _write_json(path: Path, data):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


def _slug(name: str) -> str:
    s = re.sub(r"[^a-z0-9]+", "-", (name or "").lower()).strip("-")[:40]
    return s or "campaign"


def campaigns_root() -> Path:
    d = config.campaigns_dir
    d.mkdir(parents=True, exist_ok=True)
    return d


def campaign_dir(campaign_id: str) -> Path:
    if not SAFE_ID.match(campaign_id or ""):
        raise ValueError(f"invalid campaign id: {campaign_id!r}")
    return campaigns_root() / campaign_id


def is_valid_id(campaign_id: str) -> bool:
    return bool(SAFE_ID.match(campaign_id or ""))


class Campaign:
    def __init__(self, meta: dict):
        self.meta = meta

    @property
    def id(self) -> str:
        return self.meta["id"]

    @property
    def root(self) -> Path:
        return campaign_dir(self.id)

    @property
    def input_dir(self) -> Path:
        return self.root / "input"

    @property
    def output_dir(self) -> Path:
        return self.root / "output"

    @property
    def raw_dir(self) -> Path:
        return self.root / "output" / "raw"

    @property
    def transcripts_dir(self) -> Path:
        return self.root / "transcripts"

    @property
    def context_dir(self) -> Path:
        return self.root / "context"

    @property
    def candidates_dir(self) -> Path:
        return self.root / "clip_candidates"

    @property
    def frames_dir(self) -> Path:
        return self.root / "frames"

    @property
    def meta_path(self) -> Path:
        return self.root / "meta.json"

    @property
    def rules_summary_path(self) -> Path:
        return self.root / "rules_summary.json"

    @property
    def rules_summary_legacy_path(self) -> Path:
        return self.root / "rules_summary.txt"

    @property
    def clips_path(self) -> Path:
        return self.root / "clips.json"

    @property
    def template_path(self) -> Path:
        return self.root / "template.json"

    def ensure_dirs(self):
        for d in (self.input_dir, self.output_dir, self.raw_dir,
                  self.transcripts_dir, self.context_dir, self.candidates_dir,
                  self.frames_dir):
            d.mkdir(parents=True, exist_ok=True)

    def save_meta(self):
        _write_json(self.meta_path, self.meta)

    def rules_summary(self) -> dict:
        data = _read_json(self.rules_summary_path)
        if isinstance(data, dict):
            return normalize_rules(data)
        legacy = self.rules_summary_legacy_path
        if legacy.is_file():
            try:
                text = legacy.read_text(encoding="utf-8")
            except OSError:
                text = ""
            if text.strip():
                migrated = normalize_rules({"content_criteria": _lines_to_list(text)})
                self.write_rules_summary(migrated)
                return migrated
        return empty_rules()

    def write_rules_summary(self, data):
        _write_json(self.rules_summary_path, normalize_rules(data))

    def has_rules(self) -> bool:
        r = self.rules_summary()
        if any(r.get(k) for k in RULE_LIST_SECTIONS):
            return True
        if (r.get("submission_requirements") or "").strip():
            return True
        return self.rules_full_path() is not None

    def rules_full_path(self) -> Path | None:
        for p in self.root.glob("rules_full.*"):
            if p.is_file():
                return p
        return None

    def clips(self) -> list:
        data = _read_json(self.clips_path, [])
        return data if isinstance(data, list) else []

    def save_clips(self, clips: list):
        _write_json(self.clips_path, clips)

    def clip_counts(self) -> dict:
        counts = {s: 0 for s in CLIP_STATUSES}
        for c in self.clips():
            st = c.get("status")
            if st in counts:
                counts[st] += 1
        return counts

    def has_template(self) -> bool:
        return self.template_path.is_file()

    def public(self, detail=False) -> dict:
        out = {
            "id": self.meta.get("id"),
            "name": self.meta.get("name"),
            "platform": self.meta.get("platform") or "",
            "payout_rate": self.meta.get("payout_rate") or "",
            "deadline": self.meta.get("deadline") or "",
            "status": self.meta.get("status") or "active",
            "created_at": self.meta.get("created_at") or "",
            "clip_counts": self.clip_counts(),
            "has_template": self.has_template(),
        }
        if detail:
            full = self.rules_full_path()
            out["rules_summary"] = self.rules_summary()
            out["has_rules"] = self.has_rules()
            out["rules_full"] = full.name if full else None
            out["clips"] = self.clips()
        return out


def list_campaigns() -> list[Campaign]:
    out = []
    root = campaigns_root()
    if not root.is_dir():
        return out
    for d in sorted(root.iterdir(), key=lambda p: p.name):
        if not d.is_dir():
            continue
        meta = _read_json(d / "meta.json")
        if not isinstance(meta, dict) or not meta.get("id"):
            continue
        out.append(Campaign(meta))
    out.sort(key=lambda c: c.meta.get("created_at") or "", reverse=True)
    return out


def get_campaign(campaign_id: str) -> Campaign | None:
    if not is_valid_id(campaign_id):
        return None
    meta = _read_json(campaign_dir(campaign_id) / "meta.json")
    if not isinstance(meta, dict) or not meta.get("id"):
        return None
    camp = Campaign(meta)
    camp.ensure_dirs()
    return camp


def create_campaign(name, platform="", payout_rate="", deadline="") -> Campaign:
    name = (name or "").strip()
    if not name:
        raise ValueError("name is required")
    cid = f"{_slug(name)}-{uuid.uuid4().hex[:6]}"
    meta = {
        "id": cid,
        "name": name,
        "platform": (platform or "").strip(),
        "payout_rate": (payout_rate or "").strip(),
        "deadline": (deadline or "").strip(),
        "status": "active",
        "created_at": _now(),
    }
    camp = Campaign(meta)
    camp.ensure_dirs()
    camp.save_meta()
    camp.save_clips([])
    if not camp.rules_summary_path.exists():
        camp.write_rules_summary(empty_rules())
    return camp


def update_campaign(campaign_id: str, fields: dict) -> Campaign:
    camp = get_campaign(campaign_id)
    if camp is None:
        raise FileNotFoundError(f"campaign not found: {campaign_id}")
    allowed = {"name", "platform", "payout_rate", "deadline", "status"}
    for key, val in fields.items():
        if key not in allowed:
            continue
        if key == "status":
            if val not in CAMPAIGN_STATUSES:
                raise ValueError(f"status must be one of {CAMPAIGN_STATUSES}")
            camp.meta["status"] = val
        elif key == "name":
            name = str(val or "").strip()
            if not name:
                raise ValueError("name cannot be empty")
            camp.meta["name"] = name
        else:
            camp.meta[key] = str(val or "").strip()
    camp.save_meta()
    return camp


def empty_rules() -> dict:
    return {
        "content_criteria": [],
        "brand_safety": [],
        "editing_style": [],
        "submission_requirements": "",
        "submission_done": False,
    }


def _lines_to_list(value) -> list:
    if isinstance(value, list):
        return [str(x).strip() for x in value if str(x).strip()]
    if not isinstance(value, str):
        return []
    out = []
    for ln in value.splitlines():
        item = ln.strip().lstrip("-•* ").strip()
        if item:
            out.append(item)
    return out


def normalize_rules(data) -> dict:
    base = empty_rules()
    if isinstance(data, str):
        base["content_criteria"] = _lines_to_list(data)
        return base
    if not isinstance(data, dict):
        return base
    for key in RULE_LIST_SECTIONS:
        if key in data:
            base[key] = _lines_to_list(data[key])
    if "submission_requirements" in data:
        val = data["submission_requirements"]
        if isinstance(val, list):
            base["submission_requirements"] = "\n".join(_lines_to_list(val))
        else:
            base["submission_requirements"] = str(val or "").strip()
    if "submission_done" in data:
        base["submission_done"] = bool(data["submission_done"])
    return base


def patch_rules_section(current: dict, section: str, value) -> dict:
    rules = normalize_rules(current)
    if section == "submission_done":
        rules["submission_done"] = bool(value)
        return rules
    if section not in RULE_SECTIONS:
        raise ValueError(f"unknown rules section: {section}")
    if section == "submission_requirements":
        rules[section] = str(value or "").strip() if not isinstance(value, list) \
            else "\n".join(_lines_to_list(value))
    else:
        rules[section] = _lines_to_list(value)
    return rules


def extract_rules_text(path: Path) -> str:
    """Pull plain text from a rules doc (pdf / docx / txt / md)."""
    ext = path.suffix.lower()
    if ext in {".txt", ".md"}:
        return path.read_text(encoding="utf-8", errors="replace")
    if ext == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise RuntimeError("pypdf is required to read PDF briefs") from exc
        reader = PdfReader(str(path))
        parts = []
        for page in reader.pages:
            parts.append(page.extract_text() or "")
        return "\n".join(parts)
    if ext == ".docx":
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError("python-docx is required to read Word briefs") from exc
        doc = Document(str(path))
        return "\n".join(p.text for p in doc.paragraphs)
    raise ValueError(f"unsupported rules file type: {ext}")


def summarize_rules(text: str) -> dict:
    """Condense a long brief into the four-section rules object."""
    body = (text or "").strip()
    if not body:
        return empty_rules()
    messages = [
        {"role": "system", "content": SUMMARIZE_SYSTEM},
        {"role": "user", "content": body[:24000]},
    ]
    try:
        result = call_ollama(messages, format_json=True, temperature=0.2)
        parsed = json.loads(result) if isinstance(result, str) else result
        if isinstance(parsed, str):
            parsed = json.loads(parsed)
        return normalize_rules(parsed)
    except Exception as exc:  # noqa: BLE001
        print(f"[campaigns] summarization failed: {exc}", flush=True)
        return normalize_rules({"content_criteria": _lines_to_list(body)[:15]})


def save_rules_upload(camp: Campaign, filename: str, data: bytes) -> Path:
    ext = Path(filename).suffix.lower()
    if ext not in RULES_EXTS:
        raise ValueError(f"unsupported rules type: {ext or '(none)'}")
    for old in camp.root.glob("rules_full.*"):
        if old.is_file():
            old.unlink()
    dest = camp.root / f"rules_full{ext}"
    dest.write_bytes(data)
    return dest


def upsert_clip(camp: Campaign, clip: dict) -> dict:
    clips = camp.clips()
    cid = clip.get("id")
    if cid:
        for i, existing in enumerate(clips):
            if existing.get("id") == cid:
                existing.update(clip)
                clips[i] = existing
                camp.save_clips(clips)
                return existing
    clip = dict(clip)
    clip["id"] = cid or uuid.uuid4().hex[:10]
    clips.append(clip)
    camp.save_clips(clips)
    return clip


def update_clip_status(campaign_id: str, clip_id: str, status: str) -> dict:
    if status not in CLIP_STATUSES:
        raise ValueError(f"status must be one of {CLIP_STATUSES}")
    camp = get_campaign(campaign_id)
    if camp is None:
        raise FileNotFoundError(f"campaign not found: {campaign_id}")
    clips = camp.clips()
    for c in clips:
        if c.get("id") == clip_id:
            c["status"] = status
            camp.save_clips(clips)
            return c
    raise FileNotFoundError(f"clip not found: {clip_id}")


def remove_placeholders(camp: Campaign, video_id: str):
    clips = [c for c in camp.clips()
             if not (c.get("placeholder") and c.get("video_id") == video_id)]
    camp.save_clips(clips)


def add_analyzing_placeholder(camp: Campaign, video_id: str, title=None) -> dict:
    remove_placeholders(camp, video_id)
    clip = {
        "id": f"analyzing-{video_id}",
        "video_id": video_id,
        "status": "analyzing",
        "title": title or f"Analyzing {video_id}…",
        "start": 0,
        "end": 0,
        "placeholder": True,
    }
    return upsert_clip(camp, clip)


def _clip_key(video_id, start, end):
    return (video_id, round(float(start), 3), round(float(end), 3))


def sync_clips_from_candidates(camp: Campaign, video_id: str, candidates: list,
                               status: str = "reviewing"):
    """Upsert Kanban clips from an Analyze/select candidates list."""
    remove_placeholders(camp, video_id)
    clips = camp.clips()
    index = {_clip_key(c.get("video_id"), c.get("start", 0), c.get("end", 0)): c
             for c in clips if not c.get("placeholder")}
    for cand in candidates or []:
        try:
            start, end = float(cand["start"]), float(cand["end"])
        except (KeyError, TypeError, ValueError):
            continue
        key = _clip_key(video_id, start, end)
        title = (cand.get("hook") or cand.get("reason") or "Untitled clip").strip()
        if key in index:
            rec = index[key]
            rec["title"] = title or rec.get("title")
            rec["score"] = cand.get("score", rec.get("score"))
            if rec.get("status") not in ("posted", "exported"):
                rec["status"] = status
        else:
            rec = {
                "id": uuid.uuid4().hex[:10],
                "video_id": video_id,
                "status": status,
                "title": title or "Untitled clip",
                "start": round(start, 3),
                "end": round(end, 3),
                "score": cand.get("score"),
            }
            clips.append(rec)
            index[key] = rec
    camp.save_clips(clips)
    return clips


def mark_clips_exported(camp: Campaign, video_id: str, candidates=None):
    """Advance matching clips to exported after a successful Export run."""
    clips = camp.clips()
    keys = None
    if candidates:
        keys = set()
        for cand in candidates:
            if cand.get("status") and cand.get("status") != "approved":
                continue
            try:
                keys.add(_clip_key(video_id, cand["start"], cand["end"]))
            except (KeyError, TypeError, ValueError):
                continue
    for c in clips:
        if c.get("video_id") != video_id:
            continue
        if c.get("placeholder"):
            continue
        if c.get("status") == "posted":
            continue
        if keys is not None and _clip_key(c.get("video_id"), c.get("start", 0),
                                          c.get("end", 0)) not in keys:
            continue
        if c.get("status") in ("analyzing", "reviewing", "exported"):
            c["status"] = "exported"
    camp.save_clips(clips)
    return clips
